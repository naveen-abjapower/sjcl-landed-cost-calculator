"""
Core engine for computing IEX Purchase-of-Power Landed Cost per Unit
(RTC / Peak / Normal / Off-Peak) from the two raw IEX report exports:

  1. The "SCH" file  - Scheduling Request report
     (has "Time Period" rows with "Injection/Drawal at Regional periphery"
      and "Injection/Drawal at Interface point" columns).

  2. The "DOR" file   - Daily Obligation Summary Report
     (has the Funds Payin/charges block at the top, and a
      "Daily Trade Report" table of Period / Qty in MW / Rate per MWh /
      Amount, split across two side-by-side halves of the day).

Both are plain re-implementations of the logic found in the reference
workbook "SJCL - IEX - Landed Cost - Purchase of Power.xlsx", but driven
purely off the time labels in the two source files so it keeps working
even if row positions shift from one day's report to the next.

Time-of-day buckets (Andhra Pradesh ToD schedule used in the reference
workbook):
    Peak     : 06:00-10:00  and 18:00-22:00
    Normal   : 15:00-18:00  and 22:00-24:00
    Off-Peak : 00:00-06:00  and 10:00-15:00
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field

import openpyxl

BUCKETS = ["Peak", "Normal", "OffPeak"]

# Charge line-items that the reference workbook splits EQUALLY across the
# three time buckets (flat, volume-independent fees), rather than in
# proportion to energy purchased. Matched case-insensitively as substrings
# of the charge label.
EQUAL_SPLIT_KEYWORDS = [
    "nldc application",
    "nldc scheduling",
]

TIME_RANGE_RE = re.compile(r"^\s*(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})\s*$")


class ParseError(Exception):
    """Raised when a required section/column can't be located in a file."""


# --------------------------------------------------------------------------
# Generic helpers
# --------------------------------------------------------------------------

def _minutes(h: int, m: int) -> int:
    return h * 60 + m


def bucket_for_start_label(period_label: str) -> str:
    """Classify a '06:00 - 06:15' style label into Peak/Normal/OffPeak."""
    m = TIME_RANGE_RE.match(period_label or "")
    if not m:
        raise ParseError(f"Could not parse a time range from {period_label!r}")
    h1, m1, _h2, _m2 = (int(x) for x in m.groups())
    start = _minutes(h1, m1)
    if (_minutes(6, 0) <= start < _minutes(10, 0)) or (_minutes(18, 0) <= start < _minutes(22, 0)):
        return "Peak"
    if (_minutes(15, 0) <= start < _minutes(18, 0)) or (_minutes(22, 0) <= start < _minutes(24, 0)):
        return "Normal"
    return "OffPeak"


def _is_time_label(v) -> bool:
    return isinstance(v, str) and TIME_RANGE_RE.match(v.strip()) is not None


def _to_float(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if s in ("", "-"):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


# --------------------------------------------------------------------------
# Loading workbooks (pure Python - no external programs required)
# --------------------------------------------------------------------------

def load_worksheet(file_bytes: bytes, filename: str):
    """Return an openpyxl worksheet (first sheet) with computed values."""
    suffix = os.path.splitext(filename)[1].lower()

    if suffix in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        return wb[wb.sheetnames[0]]

    if suffix == ".xls":
        return _load_xls(file_bytes)

    raise ParseError(
        f"Unsupported file type '{suffix}'. Please upload the original .xls/.xlsx export."
    )


def _load_xls(file_bytes: bytes):
    """Read a legacy .xls file with xlrd.

    Several IEX/Crystal-Reports .xls exports contain a malformed defined
    name that trips an assertion inside xlrd's name-formula evaluator
    (xlrd.book.evaluate_name_formula). We don't use defined names at all
    - only raw cell values - so we patch that one function to skip a
    name it can't parse instead of aborting the whole file load. This
    keeps everything local/offline with no external programs required.
    """
    import xlrd
    import xlrd.book as xlrd_book

    original_eval = xlrd_book.evaluate_name_formula
    if getattr(original_eval, "_landed_cost_patched", False):
        patched_eval = original_eval  # already patched by an earlier call
    else:
        def patched_eval(bk, nobj, namex, blah=False):
            try:
                return original_eval(bk, nobj, namex, blah)
            except Exception:
                return None

        patched_eval._landed_cost_patched = True
        xlrd_book.evaluate_name_formula = patched_eval

    try:
        book = xlrd.open_workbook(
            file_contents=file_bytes,
            formatting_info=False,
            ignore_workbook_corruption=True,
        )
    except Exception as e:
        raise ParseError(
            f"Couldn't read this .xls file ({e}). "
            "Please confirm it's a valid Excel export, or re-save it as "
            ".xlsx in Excel and re-upload."
        )
    finally:
        xlrd_book.evaluate_name_formula = original_eval

    return _XlrdSheetAdapter(book.sheet_by_index(0))


class _XlrdSheetAdapter:
    """Minimal adapter so xlrd sheets can be read with the same
    1-indexed ws.cell(row=, column=).value API as openpyxl."""

    def __init__(self, xlrd_sheet):
        self._sheet = xlrd_sheet
        self.max_row = xlrd_sheet.nrows
        self.max_column = xlrd_sheet.ncols

    def cell(self, row: int, column: int):
        r, c = row - 1, column - 1
        if r < 0 or c < 0 or r >= self._sheet.nrows or c >= self._sheet.ncols:
            return _Cell(None)
        return _Cell(self._sheet.cell_value(r, c))


@dataclass
class _Cell:
    value: object


# --------------------------------------------------------------------------
# SCH (Scheduling Request) file parsing -> quantities per bucket
# --------------------------------------------------------------------------

@dataclass
class ScheduleQuantities:
    """kWh per bucket, both before final losses (regional periphery /
    'net purchase') and after final losses (interface point / 'final
    purchase')."""
    net_qty_kwh: dict = field(default_factory=dict)     # before final losses
    final_qty_kwh: dict = field(default_factory=dict)   # after final losses


def parse_schedule_file(ws) -> ScheduleQuantities:
    header_row = None
    for r in range(1, min(ws.max_row, 200) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().lower() == "time period":
                header_row = r
                break
        if header_row:
            break
    if header_row is None:
        raise ParseError(
            "Couldn't find the 'Time Period' header row in the scheduling file. "
            "Is this the right file (the Scheduling Request / losses report)?"
        )

    # Find every "Time Period" column on the header row, and for each one
    # locate the nearby "Drawal at Regional periphery" / "Drawal at
    # Interface point" columns within the same block.
    period_cols = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if isinstance(v, str) and v.strip().lower() == "time period":
            period_cols.append(c)
    if not period_cols:
        raise ParseError("No 'Time Period' columns found in the scheduling file.")

    def find_col(row, start_c, end_c, needle_parts):
        for c in range(start_c, end_c + 1):
            v = ws.cell(row=row, column=c).value
            if isinstance(v, str):
                low = v.lower()
                if all(p in low for p in needle_parts):
                    return c
        return None

    blocks = []  # list of (period_col, regional_drawal_col_or_None, interface_drawal_col_or_None)
    for i, pc in enumerate(period_cols):
        next_pc = period_cols[i + 1] if i + 1 < len(period_cols) else ws.max_column + 1
        regional_col = find_col(header_row, pc, next_pc - 1, ["drawal", "regional"])
        interface_col = find_col(header_row, pc, next_pc - 1, ["drawal", "interface"])
        blocks.append((pc, regional_col, interface_col))

    net_qty = {b: 0.0 for b in BUCKETS}
    final_qty = {b: 0.0 for b in BUCKETS}

    r = header_row + 1
    rows_read = 0
    while r <= ws.max_row:
        any_time_label = False
        for pc, regional_col, interface_col in blocks:
            label = ws.cell(row=r, column=pc).value
            if not _is_time_label(label):
                continue
            any_time_label = True
            bucket = bucket_for_start_label(label)
            if regional_col:
                v = _to_float(ws.cell(row=r, column=regional_col).value) or 0.0
                net_qty[bucket] += v * 0.25 * 1000  # MW * 15min -> kWh
            if interface_col:
                v = _to_float(ws.cell(row=r, column=interface_col).value) or 0.0
                final_qty[bucket] += v * 0.25 * 1000
        if not any_time_label:
            if rows_read > 0:
                break  # ran past the data block (hit the Totals/remarks area)
        else:
            rows_read += 1
        r += 1
        if rows_read > 200:
            break

    if rows_read == 0:
        raise ParseError("No 15-minute time-block rows found under the schedule header.")

    return ScheduleQuantities(net_qty_kwh=net_qty, final_qty_kwh=final_qty)


# --------------------------------------------------------------------------
# DOR (Daily Obligation Summary Report) file parsing
# --------------------------------------------------------------------------

@dataclass
class ObligationData:
    net_amount_rs: dict = field(default_factory=dict)   # Rs per bucket, purchase amount
    net_qty_kwh: dict = field(default_factory=dict)     # kWh per bucket, from DOR's own qty col
    equal_split_charges_total_rs: float = 0.0            # flat fees, split evenly Peak/Normal/OffPeak
    proportional_charges_total_rs: float = 0.0            # fees split by energy share
    charge_lines: list = field(default_factory=list)      # [(label, amount_rs)] for transparency


def parse_obligation_file(ws) -> ObligationData:
    # ---- 1. Daily Trade Report table: Period / Qty in MW / Rate / Amount ----
    trade_header_row = None
    for r in range(1, min(ws.max_row, 300) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().lower() == "period":
                trade_header_row = r
                break
        if trade_header_row:
            break
    if trade_header_row is None:
        raise ParseError(
            "Couldn't find the 'Daily Trade Report' table (Period/Qty/Rate/Amount) "
            "in the obligation file. Is this the right file (the Daily Obligation "
            "Summary Report)?"
        )

    period_cols = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=trade_header_row, column=c).value
        if isinstance(v, str) and v.strip().lower() == "period":
            period_cols.append(c)

    def find_col(row, start_c, end_c, needle_parts):
        for c in range(start_c, end_c + 1):
            v = ws.cell(row=row, column=c).value
            if isinstance(v, str):
                low = v.lower()
                if all(p in low for p in needle_parts):
                    return c
        return None

    blocks = []
    for i, pc in enumerate(period_cols):
        next_pc = period_cols[i + 1] if i + 1 < len(period_cols) else ws.max_column + 1
        qty_col = find_col(trade_header_row, pc, next_pc - 1, ["qty"])
        amt_col = find_col(trade_header_row, pc, next_pc - 1, ["amount"])
        blocks.append((pc, qty_col, amt_col))

    net_amount = {b: 0.0 for b in BUCKETS}
    net_qty = {b: 0.0 for b in BUCKETS}

    r = trade_header_row + 1
    rows_read = 0
    while r <= ws.max_row:
        any_time_label = False
        for pc, qty_col, amt_col in blocks:
            label = ws.cell(row=r, column=pc).value
            if not _is_time_label(label):
                continue
            any_time_label = True
            bucket = bucket_for_start_label(label)
            if amt_col:
                amt = _to_float(ws.cell(row=r, column=amt_col).value) or 0.0
                net_amount[bucket] += -amt  # source amounts are negative (payin)
            if qty_col:
                qty = _to_float(ws.cell(row=r, column=qty_col).value) or 0.0
                net_qty[bucket] += qty * 0.25 * 1000
        if not any_time_label:
            if rows_read > 0:
                break
        else:
            rows_read += 1
        r += 1
        if rows_read > 200:
            break

    if rows_read == 0:
        raise ParseError("No 15-minute trade rows found in the Daily Trade Report table.")

    # ---- 2. Charges block (above the trade table): Funds Payin + line items ----
    payin_row = None
    for r in range(1, trade_header_row):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "funds payin" in v.lower():
                payin_row = r
                break
        if payin_row:
            break
    if payin_row is None:
        raise ParseError("Couldn't find the 'Funds Payin' charges section in the obligation file.")

    equal_split_total = 0.0
    proportional_total = 0.0
    charge_lines = []

    r = payin_row + 1
    while r < trade_header_row:
        label = None
        value = None
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip() and label is None:
                label = v.strip().lstrip(">").strip()
            if isinstance(v, (int, float)) and value is None:
                value = float(v)
        r += 1
        if label is None or value is None:
            continue
        low = label.lower()
        if low in ("charges", "funds payin(-) / payout(+)"):
            continue
        if low == "total":
            break  # end of the charges block
        amount = -value  # stored negative in source; convert to a positive charge
        if amount == 0:
            continue
        charge_lines.append((label, amount))
        if any(k in low for k in EQUAL_SPLIT_KEYWORDS):
            equal_split_total += amount
        else:
            proportional_total += amount

    return ObligationData(
        net_amount_rs=net_amount,
        net_qty_kwh=net_qty,
        equal_split_charges_total_rs=equal_split_total,
        proportional_charges_total_rs=proportional_total,
        charge_lines=charge_lines,
    )


# --------------------------------------------------------------------------
# Final landed-cost calculation
# --------------------------------------------------------------------------

@dataclass
class BucketResult:
    bucket: str
    final_qty_kwh: float          # A
    net_qty_kwh: float            # C (qty)
    net_amount_rs: float          # C (amount)
    other_charges_rs: float       # D
    total_amount_rs: float        # E = C + D
    trading_margin_rs: float      # F
    net_purchase_amount_rs: float  # G = E + F
    landed_cost_per_unit: float   # H = G / A


def compute_landed_cost(
    schedule: ScheduleQuantities,
    obligation: ObligationData,
    margin_rs_per_kwh: float = 0.02,
) -> dict:
    total_net_qty = sum(schedule.net_qty_kwh.values())

    results = {}
    for bucket in BUCKETS:
        A = schedule.final_qty_kwh[bucket]
        C_qty = schedule.net_qty_kwh[bucket]
        C_amt = obligation.net_amount_rs[bucket]

        share = (C_qty / total_net_qty) if total_net_qty else 0.0
        D = (obligation.equal_split_charges_total_rs / 3) + (
            obligation.proportional_charges_total_rs * share
        )
        E = C_amt + D
        F = margin_rs_per_kwh * C_qty
        G = E + F
        H = (G / A) if A else 0.0

        results[bucket] = BucketResult(
            bucket=bucket,
            final_qty_kwh=A,
            net_qty_kwh=C_qty,
            net_amount_rs=C_amt,
            other_charges_rs=D,
            total_amount_rs=E,
            trading_margin_rs=F,
            net_purchase_amount_rs=G,
            landed_cost_per_unit=H,
        )

    # RTC = simple total across all buckets (round-the-clock)
    A_rtc = sum(schedule.final_qty_kwh.values())
    C_qty_rtc = total_net_qty
    C_amt_rtc = sum(obligation.net_amount_rs.values())
    D_rtc = obligation.equal_split_charges_total_rs + obligation.proportional_charges_total_rs
    E_rtc = C_amt_rtc + D_rtc
    F_rtc = margin_rs_per_kwh * C_qty_rtc
    G_rtc = E_rtc + F_rtc
    H_rtc = (G_rtc / A_rtc) if A_rtc else 0.0

    results["RTC"] = BucketResult(
        bucket="RTC",
        final_qty_kwh=A_rtc,
        net_qty_kwh=C_qty_rtc,
        net_amount_rs=C_amt_rtc,
        other_charges_rs=D_rtc,
        total_amount_rs=E_rtc,
        trading_margin_rs=F_rtc,
        net_purchase_amount_rs=G_rtc,
        landed_cost_per_unit=H_rtc,
    )

    return results


# --------------------------------------------------------------------------
# Net Avg Cost (bottom-line margin after Andhra Pradesh Energy Cost & CSS)
# --------------------------------------------------------------------------

def compute_net_avg_cost(
    results: dict,
    ap_energy_cost: dict,
    css_rate_rs_per_kwh: float,
) -> dict:
    """Net Avg Cost per bucket = (AP Energy Cost - Landed Cost/unit) - CSS rate.

    Also returns an overall RTC figure, weighted by each bucket's final
    purchased quantity (A), matching the same weighting the reference
    workbook uses to roll bucket-level profit up to a single number.
    """
    net_avg = {}
    weighted_amount = 0.0
    weighted_qty = 0.0

    for bucket in BUCKETS:
        r = results[bucket]
        ap = ap_energy_cost.get(bucket)
        if ap is None:
            net_avg[bucket] = None
            continue
        value = (ap - r.landed_cost_per_unit) - css_rate_rs_per_kwh
        net_avg[bucket] = value
        weighted_amount += value * r.final_qty_kwh
        weighted_qty += r.final_qty_kwh

    net_avg["RTC"] = (weighted_amount / weighted_qty) if weighted_qty else None
    return net_avg
